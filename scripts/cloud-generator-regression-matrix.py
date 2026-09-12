#!/usr/bin/env python3
"""Run no-cost Cloud Generator regression scenarios in mock mode."""

from __future__ import annotations

import base64
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
PYTHONPATH = str(ROOT / "skills" / "ppt-master" / "scripts")
ACCESS_TOKEN = "matrix-token"


SCENARIOS = [
    {
        "name": "business_weekly_report",
        "brief": {"page_count": 4, "language": "zh-CN", "scenario": "executive_report", "tone": "data_driven", "include_charts": True},
        "source_text": "生成一份业务周报 PPT，包含增长、风险、下周计划。",
    },
    {
        "name": "academic_defense",
        "brief": {"page_count": 5, "language": "zh-CN", "scenario": "academic_defense", "style": "academic"},
        "source_text": "论文主题：具身智能在工业巡检中的应用。请生成答辩 PPT。",
    },
    {
        "name": "product_intro",
        "brief": {"page_count": 4, "language": "zh-CN", "scenario": "product_intro", "tone": "storytelling"},
        "source_text": "产品：AI PPT 生成器。受众：企业用户。目标：产品说明和价值展示。",
    },
    {
        "name": "market_research",
        "brief": {"page_count": 4, "language": "zh-CN", "scenario": "data_briefing", "include_charts": True},
        "source_text": "市场调研：AI 办公工具增长迅速，用户关注效率、质量和协作。",
    },
    {
        "name": "file_to_ppt",
        "brief": {"page_count": 4, "language": "zh-CN", "scenario": "general"},
        "source_text": "请根据上传材料生成 PPT。",
        "material_files": [
            {
                "name": "source-notes.txt",
                "type": "text/plain",
                "size": 80,
                "content_base64": base64.b64encode("上传材料：客户需求、竞品分析、路线图。".encode("utf-8")).decode("ascii"),
            }
        ],
    },
    {
        "name": "template_parameter",
        "brief": {"page_count": 4, "language": "zh-CN", "template_id": "matrix-smoke-template", "template_kind": "deck"},
        "source_text": "用指定模板参数生成一个项目复盘 PPT。",
    },
    {
        "name": "image_generation_option",
        "brief": {"page_count": 4, "language": "zh-CN", "include_images": True, "image_source_mode": "generate"},
        "source_text": "生成一份城市低空经济科普 PPT，需要配图建议。",
        "expect_images": True,
    },
    {
        "name": "audio_export_option",
        "brief": {"page_count": 4, "language": "zh-CN", "include_audio": True, "audio_provider": "edge"},
        "source_text": "生成一份带讲稿旁白的培训 PPT。",
        "expect_audio": True,
    },
    {
        "name": "audio_animation_combination",
        "brief": {"page_count": 4, "language": "zh-CN", "include_audio": True, "include_animations": True, "audio_provider": "edge"},
        "source_text": "生成一份带讲稿旁白和简单动画的培训 PPT。",
        "expect_audio": True,
    },
]


def main() -> int:
    jobs_dir = Path(tempfile.mkdtemp(prefix="ppt-master-cloud-matrix-"))
    port = _free_port()
    server = _start_server(port, jobs_dir)
    base = f"http://127.0.0.1:{port}"
    results = []
    try:
        _wait_health(base)
        _assert_route_contract(base)
        session = _register(base)
        _assert_unsupported_route_rejected(base, session)
        _json("POST", base, "/billing/checkout", {"plan": "team", "auto_complete": True}, session=session)
        _assert_fill_native_route_session(base, session)
        for scenario in SCENARIOS:
            task_id = _create_task(base, session, scenario)
            task = _wait_task(base, task_id, session)
            if task.get("status") != "completed":
                raise AssertionError(f"{scenario['name']} failed: {task.get('error_message')}")
            preview = _json("GET", base, f"/generation-tasks/{task_id}/preview", session=session)
            artifacts = _json("GET", base, f"/generation-tasks/{task_id}/artifacts", session=session)
            if scenario.get("expect_images") and not preview.get("images"):
                raise AssertionError(f"{scenario['name']} expected image assets")
            if scenario.get("expect_audio") and not any(item.get("kind") == "audio" for item in artifacts.get("artifacts") or []):
                raise AssertionError(f"{scenario['name']} expected audio artifacts")
            download = _json("GET", base, f"/generation-tasks/{task_id}/download", session=session)
            if scenario.get("expect_audio"):
                if not download.get("narrated_ready") or not download.get("narrated_download_url"):
                    raise AssertionError(f"{scenario['name']} expected narrated PPTX download metadata")
                if not any(item.get("kind") == "narrated_pptx" for item in artifacts.get("artifacts") or []):
                    raise AssertionError(f"{scenario['name']} expected narrated PPTX artifact")
            elif download.get("narrated_ready"):
                raise AssertionError(f"{scenario['name']} should not expose narrated PPTX")
            metrics = task.get("metrics") or {}
            results.append({
                "name": scenario["name"],
                "task_id": task_id,
                "pages": metrics.get("page_count_generated"),
                "attempts": metrics.get("page_attempts"),
                "artifacts": len(artifacts.get("artifacts") or []),
                "images": len(preview.get("images") or []),
                "audio": sum(1 for item in artifacts.get("artifacts") or [] if item.get("kind") == "audio"),
            })
        print(json.dumps({"ok": True, "scenarios": results}, ensure_ascii=False, indent=2))
        return 0
    finally:
        server.terminate()
        try:
            server.wait(timeout=5)
        except subprocess.TimeoutExpired:
            server.kill()


def _start_server(port: int, jobs_dir: Path) -> subprocess.Popen:
    env = os.environ.copy()
    env["CLOUD_GENERATOR_MOCK"] = "1"
    env["PPT_MASTER_REQUIRE_AUTH"] = "1"
    env["PYTHONPATH"] = PYTHONPATH
    return subprocess.Popen(
        [
            sys.executable,
            "-m",
            "cloud_generator.api_server",
            "--host",
            "127.0.0.1",
            "--port",
            str(port),
            "--store",
            "json",
            "--jobs-dir",
            str(jobs_dir),
            "--access-token",
            ACCESS_TOKEN,
            "--frontend-base-url",
            "http://127.0.0.1:8000",
        ],
        cwd=str(ROOT),
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def _register(base: str) -> str:
    payload = {"email": f"matrix-{uuid.uuid4().hex[:12]}@example.com", "password": "pptmaster", "auto_register": True}
    data = _json("POST", base, "/auth/login", payload, include_access=False)
    return data["session_token"]


def _assert_route_contract(base: str) -> None:
    capabilities = _json("GET", base, "/core-capabilities")
    rows = {item.get("route"): item for item in capabilities.get("capabilities") or []}
    if rows.get("generate_pptx", {}).get("available") is not True:
        raise AssertionError("generate route should be cloud-enabled")
    for route in ("create_template", "fill_native_pptx", "enhance_native_pptx"):
        if rows.get(route, {}).get("core_available") is not True:
            raise AssertionError(f"{route} should be present in core")
        if rows.get(route, {}).get("available") is not False:
            raise AssertionError(f"{route} should be gated in cloud")


def _assert_unsupported_route_rejected(base: str, session: str) -> None:
    status, data = _request(
        "POST",
        base,
        "/generation-tasks",
        {
            "job_id": f"matrix_unsupported_{uuid.uuid4().hex[:8]}",
            "run_async": True,
            "brief": {"route": "fill_native_pptx", "page_count": 4},
            "source_text": "unsupported route preflight",
        },
        session=session,
    )
    if status != 400 or "not available" not in str(data.get("error") or ""):
        raise AssertionError(f"unsupported route should fail preflight: HTTP {status}: {data}")


def _assert_fill_native_route_session(base: str, session: str) -> None:
    from pptx import Presentation

    before_quota = _json("GET", base, "/membership", session=session).get("quota", {})

    with tempfile.TemporaryDirectory(prefix="ppt-master-native-route-fixture-") as temp_dir:
        source = Path(temp_dir) / "native-template.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(source)
        material = {
            "name": source.name,
            "type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "size": source.stat().st_size,
            "content_base64": base64.b64encode(source.read_bytes()).decode("ascii"),
        }
        prepared = _json(
            "POST",
            base,
            "/route-sessions",
            {
                "route": "fill_native_pptx",
                "source_text": "把这页填入新的项目摘要。",
                "material_files": [material],
            },
            session=session,
        )
        if prepared.get("status") != "awaiting_confirmation":
            raise AssertionError(f"Fill Native should stop for confirmation: {prepared}")
        state = _json("GET", base, prepared["session_url"], session=session)
        plan = state.get("confirmation", {}).get("fill_plan")
        if not isinstance(plan, dict):
            raise AssertionError(f"Fill Native response omitted the confirmation plan: {state}")
        plan["user_reviewed"] = True
        completed = _json(
            "POST",
            base,
            f"/route-sessions/{prepared['session_id']}/confirm",
            {"confirmed": True, "idempotency_key": f"confirm:{prepared['session_id']}", "fill_plan": plan},
            session=session,
        )
        download = _json("GET", base, completed["download_url"], session=session)
        if completed.get("status") != "completed" or not download.get("download_url", "").endswith("/download-file"):
            raise AssertionError(f"Fill Native confirmation did not produce an output: {completed}")
        repeated = _json(
            "POST",
            base,
            f"/route-sessions/{prepared['session_id']}/confirm",
            {"confirmed": True, "idempotency_key": f"confirm:{prepared['session_id']}", "fill_plan": plan},
            session=session,
        )
        if repeated.get("status") != "completed":
            raise AssertionError(f"Repeated Fill Native confirmation was not idempotent: {repeated}")

        template = _json(
            "POST",
            base,
            "/route-sessions",
            {
                "route": "create_template",
                "source_text": "发布一套可复用的品牌模板。",
                "material_files": [material],
            },
            session=session,
        )
        if template.get("status") != "awaiting_authoring":
            raise AssertionError(f"Create Template should expose authoring gate: {template}")
        published = _json(
            "POST",
            base,
            f"/route-sessions/{template['session_id']}/publish",
            {"name": "Matrix reusable template", "description": "Regression fixture", "authoring": {"primary_color": "#123456", "font_family": "Aptos", "tags": ["regression"]}, "idempotency_key": f"publish:{template['session_id']}"},
            session=session,
        )
        if published.get("status") != "published" or not published.get("template", {}).get("template_id"):
            raise AssertionError(f"Create Template publish did not complete: {published}")
        templates = _json("GET", base, "/templates", session=session).get("templates") or []
        if not any(item.get("id") == published["template"]["template_id"] for item in templates):
            raise AssertionError("Published template was not returned from /templates")
        repeated_publish = _json(
            "POST",
            base,
            f"/route-sessions/{template['session_id']}/publish",
            {"name": "Matrix reusable template", "idempotency_key": f"publish:{template['session_id']}"},
            session=session,
        )
        if repeated_publish.get("status") != "published":
            raise AssertionError(f"Repeated template publish was not idempotent: {repeated_publish}")
        after_quota = _json("GET", base, "/membership", session=session).get("quota", {})
        used_delta = int(after_quota.get("used") or 0) - int(before_quota.get("used") or 0)
        if used_delta <= 0:
            raise AssertionError(f"Native route sessions did not debit quota: before={before_quota} after={after_quota}")
        template_id = published["template"]["template_id"]
        generated = _json(
            "POST",
            base,
            "/generation-tasks",
            {
                "job_id": f"matrix_published_template_{uuid.uuid4().hex[:8]}",
                "run_async": True,
                "brief": {
                    "page_count": 4,
                    "language": "zh-CN",
                    "template_id": template_id,
                    "template_kind": "deck",
                    "template_path": next(item["path"] for item in templates if item.get("id") == template_id),
                    "source_text": "使用已发布模板生成一份项目摘要。",
                },
                "source_text": "使用已发布模板生成一份项目摘要。",
            },
            session=session,
        )
        generated_state = _wait_task(base, generated["task_id"], session)
        if generated_state.get("status") != "completed":
            raise AssertionError(f"Published template could not be consumed by Generate: {generated_state}")


def _create_task(base: str, session: str, scenario: dict) -> str:
    task_id = f"matrix_{scenario['name']}_{int(time.time() * 1000)}"
    brief = dict(scenario["brief"], source_text=scenario["source_text"])
    if scenario["name"] == "template_parameter":
        templates = _json("GET", base, "/templates", session=session).get("templates") or []
        if templates:
            selected = templates[0]
            brief.update({
                "template_id": selected.get("id") or "",
                "template_kind": selected.get("kind") or "",
                "template_path": selected.get("path") or "",
            })
    payload = {
        "job_id": task_id,
        "run_async": True,
        "max_attempts": 2,
        "brief": brief,
        "source_text": scenario["source_text"],
        "material_files": scenario.get("material_files") or [],
    }
    data = _json("POST", base, "/generation-tasks", payload, session=session)
    return data["task_id"]


def _wait_task(base: str, task_id: str, session: str) -> dict:
    deadline = time.time() + 30
    last = {}
    while time.time() < deadline:
        last = _json("GET", base, f"/generation-tasks/{task_id}", session=session)
        if last.get("status") in {"completed", "failed", "cancelled"}:
            return last
        time.sleep(0.25)
    raise RuntimeError(f"task did not finish: {last}")


def _wait_health(base: str) -> None:
    deadline = time.time() + 10
    while time.time() < deadline:
        try:
            if _json("GET", base, "/health").get("status") == "ok":
                return
        except Exception:
            time.sleep(0.2)
    raise RuntimeError("API server did not become healthy")


def _json(method: str, base: str, path: str, payload: dict | None = None, *, session: str = "", include_access: bool = True) -> dict:
    status, data = _request(method, base, path, payload, session=session, include_access=include_access)
    if not (200 <= status < 300):
        raise RuntimeError(f"{method} {path} failed: HTTP {status}: {data}")
    return data


def _request(method: str, base: str, path: str, payload: dict | None = None, *, session: str = "", include_access: bool = True) -> tuple[int, dict]:
    headers = {}
    body = None
    if include_access:
        headers["X-PPT-Master-Token"] = ACCESS_TOKEN
    if session:
        headers["X-PPT-Master-Session"] = session
    if payload is not None:
        headers["Content-Type"] = "application/json"
        body = json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(base + path, data=body, method=method, headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8") or "{}")
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode("utf-8", errors="replace")
        try:
            data = json.loads(raw or "{}")
        except Exception:
            data = {"raw": raw}
        return exc.code, data


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


if __name__ == "__main__":
    raise SystemExit(main())
